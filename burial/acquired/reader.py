# -*- coding: utf-8 -*-
"""Delivered tables to numpy columns: CSV / text / XLSX, read by an import profile.

Two steps, so the import screen can preview a file before it has a profile:

1. :func:`read_table` splits the delivery into a header, an optional units
   line and data rows, each with its source line number. :func:`sniff`
   proposes the table settings (delimiter, header and units lines, decimal
   comma, text encoding).
2. :func:`parse` applies a profile: each mapped column is read in its role's
   way (numbers with units and sign, timestamps to UTC seconds, coordinates
   in decimal or degrees-minutes) and converted to the stored units.

A row is kept or rejected as a whole, with a plain reason ("DOL: '#VALUE!'
is not a number", "no KP", "row ends early: 9 of 14 values"); missing-value
codes (blank, ``n/a``, ``-9999``…) are not errors but NaN. Columns the
profile leaves unmapped are kept as extra channels when they are numeric,
leniently (a stray text cell there becomes NaN, never a rejection).

Standard library + numpy only (XLSX through the bundled openpyxl, imported
when an XLSX is read).
"""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np

from . import profile as P
from . import roles as R

_EPOCH = datetime(1970, 1, 1)
_SNIFF_LINES = 60
_EXTRA_NUMERIC_SHARE = 0.9

DATE_PATTERNS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%Y", "%Y%m%d",
                 "%Y:%j", "%d-%b-%Y", "%d %b %Y", "%d/%m/%y", "%m/%d/%y")
CLOCK_PATTERNS = ("%H:%M:%S", "%H:%M", "%H%M%S")
TIMESTAMP_PATTERNS = tuple(f"{d} {c}" for d in DATE_PATTERNS for c in CLOCK_PATTERNS[:2]) + (
    "%Y:%j:%H:%M:%S", "%d/%m/%Y,%H:%M:%S", "%Y%m%d%H%M%S", "%Y%m%dT%H%M%S")


class ReadError(ValueError):
    """The delivery cannot be read as a table (message is user-facing)."""


@dataclass
class RawTable:
    headers: List[str]
    units: List[str]
    rows: List[List[str]]
    lines: List[int]                 # source line (text) or sheet row (XLSX) of each row
    delimiter: str = ","
    encoding: str = "utf-8"
    kind: str = "text"               # "text" | "xlsx"
    sheet: Optional[str] = None
    sheets: List[str] = field(default_factory=list)

    def row_text(self, index: int) -> str:
        sep = " " if self.delimiter == "whitespace" else self.delimiter
        return sep.join(self.rows[index])


@dataclass
class Channel:
    key: str
    label: str
    role: str
    unit: str
    header: str
    auto: bool = False               # an extra channel kept because it is numeric


@dataclass
class ParsedData:
    columns: Dict[str, np.ndarray]   # float64 per channel key, stored units
    texts: Dict[str, List[str]]      # event / pass text per kept row
    channels: List[Channel]
    line: np.ndarray                 # int64 source line of each kept row
    rejected: List[Tuple[int, str, str]]   # (line, reason, raw text)
    warnings: List[str]
    rows_read: int
    time_format: str = ""

    @property
    def n(self) -> int:
        return int(self.line.size)

    def _range(self, key: str) -> Tuple[Optional[float], Optional[float]]:
        values = self.columns.get(key)
        if values is None or not values.size or not np.isfinite(values).any():
            return None, None
        return float(np.nanmin(values)), float(np.nanmax(values))

    @property
    def kp_range(self) -> Tuple[Optional[float], Optional[float]]:
        return self._range(R.KP)

    @property
    def time_range(self) -> Tuple[Optional[float], Optional[float]]:
        return self._range(R.TIME)

    def channel(self, key: str) -> Optional[Channel]:
        return next((c for c in self.channels if c.key == key), None)


# --------------------------------------------------------------------------- reading the table
def is_xlsx(name: str, content: bytes = b"") -> bool:
    return str(name).lower().endswith((".xlsx", ".xlsm")) or content[:4] == b"PK\x03\x04"


def decode(content: bytes, encoding: Optional[str] = None) -> Tuple[str, str]:
    """Text and the encoding used: UTF-8 (with or without BOM), else Windows-1252."""
    if encoding:
        return content.decode(encoding), encoding
    try:
        return content.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace"), "cp1252"


def _split(line: str, delimiter: str) -> List[str]:
    if delimiter == "whitespace":
        return line.split()
    return next(csv.reader([line], delimiter=delimiter), [])


def _guess_delimiter(lines: Sequence[str]) -> str:
    best, best_score = None, 0.0
    for candidate in (",", ";", "\t", "|"):
        counts = [line.count(candidate) for line in lines]
        hits = [c for c in counts if c]
        if len(hits) < max(1, len(lines) // 2):
            continue
        mode = max(set(hits), key=hits.count)
        score = hits.count(mode) / len(lines) * (1 + min(mode, 30) / 30)
        if score > best_score:
            best, best_score = candidate, score
    if best:
        return best
    if lines and sum(len(line.split()) > 1 for line in lines) >= len(lines) / 2:
        return "whitespace"
    return ","


def _numeric(cell: str, decimal_comma: bool = False) -> bool:
    text = cell.strip()
    if decimal_comma:
        text = text.replace(",", ".")
    try:
        float(text)
        return True
    except ValueError:
        return False


def _mostly_numeric(cells: Sequence[str], decimal_comma: bool) -> bool:
    filled = [c for c in cells if c.strip()]
    return bool(filled) and sum(_numeric(c, decimal_comma) for c in filled) >= 0.6 * len(filled)


def _mostly_text(cells: Sequence[str]) -> bool:
    filled = [c for c in cells if c.strip()]
    return len(filled) >= 2 and sum(not _numeric(c, True) for c in filled) >= 0.6 * len(filled)


def sniff(content: bytes, name: str = "") -> Dict:
    """Proposed table settings for a delivery (``profile["table"]`` keys)."""
    table = dict(P.new_profile()["table"])
    if is_xlsx(name, content):
        sample = _xlsx_rows(content, None, _SNIFF_LINES)[0]
        lines = [(number, [_cell_text(c) for c in cells]) for number, cells in sample
                 if any(c is not None for c in cells)]
        table["delimiter"] = None
    else:
        text, encoding = decode(content[:262144])
        table["encoding"] = None if encoding == "utf-8" else encoding
        raw_lines = [line for line in text.splitlines()[:_SNIFF_LINES]]
        filled = [line for line in raw_lines if line.strip()]
        delimiter = _guess_delimiter(filled[:30])
        table["delimiter"] = delimiter
        lines = [(i + 1, _split(line, delimiter)) for i, line in enumerate(raw_lines) if line.strip()]
    if not lines:
        return table
    comma = table.get("delimiter") not in (",", None) and _decimal_comma_share(lines) > 0.3
    table["decimal_comma"] = bool(comma)
    header_at = 0
    for index, (_number, cells) in enumerate(lines):
        following = [c for _n, c in lines[index + 1:index + 4]]
        if _mostly_text(cells) and any(_mostly_numeric(f, comma) for f in following):
            header_at = index
            break
    table["header_line"] = lines[header_at][0]
    if header_at + 2 < len(lines):
        units_cells = lines[header_at + 1][1]
        after = lines[header_at + 2][1]
        if (not _mostly_numeric(units_cells, comma) and _mostly_numeric(after, comma)
                and all(len(c.strip()) <= 14 for c in units_cells)):
            table["units_line"] = lines[header_at + 1][0]
    return table


def _decimal_comma_share(lines) -> float:
    pattern = re.compile(r"^[+-]?\d+,\d+$")
    cells = [c.strip() for _n, row in lines for c in row if c.strip()]
    numeric = [c for c in cells if _numeric(c, True)]
    return sum(bool(pattern.match(c)) for c in numeric) / len(numeric) if numeric else 0.0


def _unique_headers(headers: Sequence[str]) -> List[str]:
    out, seen = [], {}
    for index, header in enumerate(headers):
        name = str(header).strip() or f"Column {index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def read_table(content: bytes, name: str, table: Optional[Dict] = None,
               max_rows: Optional[int] = None) -> RawTable:
    """Header, units line and data rows of a delivery (rows may be ragged)."""
    settings = dict(P.new_profile()["table"])
    settings.update(table or sniff(content, name))
    header_line = max(1, int(settings.get("header_line") or 1))
    units_line = int(settings.get("units_line") or 0)
    first_data = max(header_line, units_line) + 1
    if is_xlsx(name, content):
        rows, sheet, sheets = _xlsx_rows(content, settings.get("sheet"), None if max_rows is None
                                         else first_data + max_rows)
        by_line = {number: [_cell_text(c) for c in cells] for number, cells in rows}
        if header_line not in by_line:
            raise ReadError(f"Row {header_line} of sheet '{sheet}' is empty: choose the header row.")
        raw = RawTable(_unique_headers(_trim(by_line[header_line], 0)),
                       by_line.get(units_line, []) if units_line else [],
                       [], [], delimiter="\t", kind="xlsx", sheet=sheet, sheets=sheets)
        width = len(raw.headers)
        for number, cells in sorted(by_line.items()):
            if number >= first_data and any(c.strip() for c in cells):
                # A worksheet row has every cell: blanks at its end are empty cells, not a short row.
                cells = _trim(cells, width)
                raw.rows.append(cells + [""] * (width - len(cells)))
                raw.lines.append(number)
        return raw
    text, encoding = decode(content, settings.get("encoding"))
    delimiter = settings.get("delimiter") or _guess_delimiter(
        [line for line in text.splitlines()[:30] if line.strip()])
    lines = text.splitlines()
    if header_line > len(lines):
        raise ReadError(f"The file has only {len(lines)} lines; the header cannot be on line {header_line}.")
    headers = _unique_headers(_split(lines[header_line - 1], delimiter))
    if len(headers) < 2:
        raise ReadError(f"Line {header_line} does not split into columns with this delimiter.")
    units = _split(lines[units_line - 1], delimiter) if 0 < units_line <= len(lines) else []
    raw = RawTable(headers, units, [], [], delimiter=delimiter, encoding=encoding)
    width = len(headers)
    body = lines[first_data - 1:]
    if delimiter == "whitespace":
        for offset, line in enumerate(body):
            if line.strip():
                raw.rows.append(line.split())
                raw.lines.append(first_data + offset)
                if max_rows is not None and len(raw.rows) >= max_rows:
                    break
    else:
        reader = csv.reader(io.StringIO("\n".join(body)), delimiter=delimiter)
        start = 1
        for cells in reader:
            line_number = first_data + start - 1
            start = reader.line_num + 1
            if any(c.strip() for c in cells):
                raw.rows.append(_trim(cells, width))
                raw.lines.append(line_number)
                if max_rows is not None and len(raw.rows) >= max_rows:
                    break
    return raw


def _trim(cells: List[str], width: int) -> List[str]:
    """Drop empty cells past ``width`` (a trailing delimiter); a short row stays short."""
    cells = list(cells)
    while len(cells) > width and not str(cells[-1]).strip():
        cells.pop()
    return cells


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    return str(value)


def _xlsx_rows(content: bytes, sheet: Optional[str], limit: Optional[int]):
    """``([(row number, values)], sheet name, sheet names)`` of one worksheet."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - openpyxl is bundled in lib/
        raise ReadError("Reading Excel files needs openpyxl, which could not be loaded.") from exc
    try:
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ReadError(f"Could not open the workbook: {exc}") from exc
    try:
        names = list(book.sheetnames)
        name = sheet if sheet in names else names[0]
        rows = []
        for number, values in enumerate(book[name].iter_rows(values_only=True), start=1):
            rows.append((number, list(values)))
            if limit is not None and number >= limit:
                break
        return rows, name, names
    finally:
        book.close()


# --------------------------------------------------------------------------- values
def parse_numbers(cells: Sequence[str], decimal_comma: bool = False,
                  missing: Sequence[str] = P.DEFAULT_MISSING) -> Tuple[np.ndarray, np.ndarray]:
    """``(values, bad)``: float64 with NaN for missing codes; ``bad`` marks text that is no number."""
    n = len(cells)
    codes = {str(m).strip().lower() for m in missing}
    numeric_codes = {float(m) for m in codes if _numeric(m)}
    if not decimal_comma:
        try:
            values = np.array([c if str(c).strip() else "nan" for c in cells], dtype=np.float64)
        except ValueError:
            values = None
        if values is not None:
            bad = np.isinf(values)
            values[bad] = np.nan
            if numeric_codes:
                values[np.isin(values, list(numeric_codes))] = np.nan
            return values, bad
    values = np.full(n, np.nan)
    bad = np.zeros(n, dtype=bool)
    for i, cell in enumerate(cells):
        text = str(cell).strip()
        if text.lower() in codes:
            continue
        if decimal_comma:
            text = text.replace(",", ".")
        try:
            value = float(text)
        except ValueError:
            bad[i] = True
            continue
        if not math.isfinite(value):
            if math.isinf(value):
                bad[i] = True
            continue
        if value not in numeric_codes:
            values[i] = value
    return values, bad


_HEMI = re.compile(r"[NSEW]")


def parse_coordinate(text: str, decimal_comma: bool = False) -> float:
    """Degrees from ``55.123``, ``55°07.38'N``, ``N55 07 22.8``, ``003-07.12W``; NaN if not a coordinate."""
    value = str(text or "").strip().upper()
    if not value:
        return math.nan
    if decimal_comma:
        value = value.replace(",", ".")
    try:
        return float(value)
    except ValueError:
        pass
    sign = 1.0
    hemi = _HEMI.findall(value)
    if len(hemi) > 1:
        return math.nan
    if hemi and hemi[0] in "SW":
        sign = -1.0
    body = _HEMI.sub(" ", value).strip()
    if body.startswith("-"):
        sign, body = -sign, body[1:]
    body = body.lstrip("+")
    parts = [p for p in re.split(r"[°º'′\"″:\s]+|(?<=\d)-(?=\d)", body) if p]
    try:
        numbers = [float(p) for p in parts]
    except ValueError:
        return math.nan
    if not 1 <= len(numbers) <= 3 or any(v < 0 for v in numbers) or any(v >= 60 for v in numbers[1:]):
        return math.nan
    degrees = numbers[0] + (numbers[1] / 60.0 if len(numbers) > 1 else 0.0) + (
        numbers[2] / 3600.0 if len(numbers) > 2 else 0.0)
    return sign * degrees


def parse_coordinates(cells: Sequence[str], decimal_comma: bool, missing: Sequence[str],
                      limit: float) -> Tuple[np.ndarray, np.ndarray]:
    values, bad = parse_numbers(cells, decimal_comma, missing)
    if bad.any():
        for i in np.flatnonzero(bad):
            value = parse_coordinate(cells[i], decimal_comma)
            if math.isfinite(value):
                values[i], bad[i] = value, False
    out_of_range = np.isfinite(values) & (np.abs(values) > limit)
    values[out_of_range] = np.nan
    return values, bad | out_of_range


def _naive_seconds(dt: datetime) -> float:
    return (dt - _EPOCH).total_seconds()


def _to_utc_seconds(dt: datetime, offset_s: int) -> float:
    if dt.tzinfo is not None:
        return _naive_seconds(dt.astimezone(timezone.utc).replace(tzinfo=None))
    return _naive_seconds(dt) - offset_s


def _fits(pattern: str, samples: Sequence[str]) -> bool:
    for text in samples:
        if _strptime(text, pattern) is None:
            return False
    return True


def _strptime(text: str, pattern: str) -> Optional[datetime]:
    try:
        return datetime.strptime(text, pattern)
    except ValueError:
        if "." in text and "%S" in pattern and "%f" not in pattern:
            try:
                return datetime.strptime(text, pattern.replace("%S", "%S.%f"))
            except ValueError:
                return None
        return None


def _iso(text: str) -> Optional[datetime]:
    try:
        return datetime.fromisoformat(text.replace("z", "Z"))
    except ValueError:
        return None


def _choose(patterns: Sequence[str], samples: Sequence[str]) -> Tuple[Optional[str], bool]:
    """First pattern that reads every sample, and whether day/month order was ambiguous."""
    fitting = [p for p in patterns if _fits(p, samples)]
    if not fitting:
        return None, False
    first = fitting[0]
    swapped = first.replace("%d", "\0").replace("%m", "%d").replace("\0", "%m")
    return first, swapped != first and swapped in fitting


def detect_time_format(cells: Sequence[str]) -> Tuple[Optional[str], bool]:
    """``(format, ambiguous)`` for a timestamp column: ``iso``, ``epoch``, ``epoch_ms`` or a pattern."""
    samples = _samples(cells)
    if not samples:
        return None, False
    try:
        numbers = [float(s) for s in samples]
        if all(n >= 1e11 for n in numbers):
            return "epoch_ms", False
        if all(n >= 1e8 for n in numbers) and not all(re.fullmatch(r"\d{8}", s) for s in samples):
            return "epoch", False
    except ValueError:
        pass
    if all(re.match(r"\d{4}-\d{2}-\d{2}", s) and _iso(s) is not None for s in samples):
        return "iso", False
    return _choose(TIMESTAMP_PATTERNS, samples)


def detect_date_format(cells: Sequence[str]) -> Tuple[Optional[str], bool]:
    samples = _samples(cells)
    if samples and all(re.fullmatch(r"\d{4}-\d{2}-\d{2}(T00:00:00)?", s) for s in samples):
        return "%Y-%m-%d", False
    return _choose(DATE_PATTERNS, samples)


def _samples(cells: Sequence[str], count: int = 200) -> List[str]:
    filled = [str(c).strip() for c in cells if str(c).strip()]
    if len(filled) <= count:
        return filled
    step = len(filled) / count
    return [filled[int(i * step)] for i in range(count)]


def parse_timestamps(cells: Sequence[str], fmt: str, offset_s: int,
                     missing: Sequence[str] = P.DEFAULT_MISSING) -> Tuple[np.ndarray, np.ndarray, str, bool]:
    """``(UTC seconds, bad, format used, day/month ambiguous)``; times carrying an offset keep it."""
    codes = {str(m).strip().lower() for m in missing}
    ambiguous = False
    if fmt in (None, "", "auto"):
        fmt, ambiguous = detect_time_format(cells)
        if fmt is None:
            return np.full(len(cells), np.nan), np.array([bool(str(c).strip()) for c in cells]), "", False
    n = len(cells)
    values = np.full(n, np.nan)
    bad = np.zeros(n, dtype=bool)
    if fmt in ("epoch", "epoch_ms"):
        numbers, bad = parse_numbers(cells, False, missing)
        return (numbers / 1000.0 if fmt == "epoch_ms" else numbers), bad, fmt, False
    for i, cell in enumerate(cells):
        text = str(cell).strip()
        if text.lower() in codes:
            continue
        dt = _iso(text) if fmt == "iso" else _strptime(text, fmt)
        if dt is None:
            bad[i] = True
        else:
            values[i] = _to_utc_seconds(dt, offset_s)
    return values, bad, fmt, ambiguous


def parse_clock(cells: Sequence[str], missing: Sequence[str] = P.DEFAULT_MISSING) -> Tuple[np.ndarray, np.ndarray]:
    """Seconds after midnight from ``HH:MM[:SS[.f]]`` or ``HHMMSS``."""
    codes = {str(m).strip().lower() for m in missing}
    values = np.full(len(cells), np.nan)
    bad = np.zeros(len(cells), dtype=bool)
    for i, cell in enumerate(cells):
        text = str(cell).strip()
        if text.lower() in codes:
            continue
        if "T" in text:      # an XLSX time cell read as a full datetime
            text = text.split("T", 1)[1]
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}(?:\.\d*)?))?", text) or \
            re.fullmatch(r"(\d{2})(\d{2})(\d{2}(?:\.\d*)?)", text)
        if not match:
            bad[i] = True
            continue
        hours, minutes = int(match.group(1)), int(match.group(2))
        seconds = float(match.group(3) or 0)
        if hours > 24 or minutes > 59 or seconds >= 61:
            bad[i] = True
            continue
        values[i] = hours * 3600 + minutes * 60 + seconds
    return values, bad


def parse_dates(cells: Sequence[str], fmt: str, missing: Sequence[str] = P.DEFAULT_MISSING
                ) -> Tuple[np.ndarray, np.ndarray, str, bool]:
    """``(seconds at 00:00 of each date, bad, format, ambiguous)``; repeated dates parsed once."""
    codes = {str(m).strip().lower() for m in missing}
    ambiguous = False
    if fmt in (None, "", "auto"):
        fmt, ambiguous = detect_date_format(cells)
        if fmt is None:
            return np.full(len(cells), np.nan), np.array([bool(str(c).strip()) for c in cells]), "", False
    cache: Dict[str, Optional[float]] = {}
    values = np.full(len(cells), np.nan)
    bad = np.zeros(len(cells), dtype=bool)
    for i, cell in enumerate(cells):
        text = str(cell).strip()
        if text.lower() in codes:
            continue
        if text not in cache:
            dt = _strptime(text.split("T", 1)[0] if fmt == "%Y-%m-%d" else text, fmt)
            cache[text] = None if dt is None else _naive_seconds(datetime(dt.year, dt.month, dt.day))
        seconds = cache[text]
        if seconds is None:
            bad[i] = True
        else:
            values[i] = seconds
    return values, bad, fmt, ambiguous


# --------------------------------------------------------------------------- applying a profile
def parse(raw: RawTable, profile: Dict, cancelled: Optional[Callable[[], bool]] = None) -> ParsedData:
    """Read every row of ``raw`` through ``profile`` (see the module notes)."""
    prof = P.normalise(profile)
    problems = P.problems(prof, raw.headers)
    if problems:
        raise ReadError(" ".join(problems))
    table = prof["table"]
    comma = bool(table.get("decimal_comma"))
    missing = table.get("missing") or P.DEFAULT_MISSING
    offset = P.parse_utc_offset(prof.get("timezone") or "UTC") or 0
    headers = raw.headers
    width = len(headers)
    index = {h: i for i, h in enumerate(headers)}
    n = len(raw.rows)
    reasons: List[Optional[str]] = [None] * n

    def reject(mask: np.ndarray, reason: Callable[[int], str]) -> None:
        for i in np.flatnonzero(mask):
            if reasons[i] is None:
                reasons[i] = reason(int(i))

    entries = P.mapped(prof)
    strict_cols = {index[e["header"]] for e in entries if R.role(e["role"]).kind != R.KIND_TEXT}
    lengths = [len(r) for r in raw.rows]
    for i, length in enumerate(lengths):
        if length > width:
            reasons[i] = f"{length} values where the header has {width}"
        elif length < width and any(c >= length for c in strict_cols):
            reasons[i] = f"row ends early: {length} of {width} values"

    def cells_of(col: int) -> List[str]:
        return [row[col] if col < len(row) else "" for row in raw.rows]

    columns: Dict[str, np.ndarray] = {}
    texts: Dict[str, List[str]] = {}
    channels: List[Channel] = []
    warnings: List[str] = []
    time_parts: Dict[str, np.ndarray] = {}
    time_format = ""
    for entry in entries:
        if cancelled and cancelled():
            raise InterruptedError("Import cancelled")
        role = R.role(entry["role"])
        header = entry["header"]
        cells = cells_of(index[header])
        label = role.label
        if role.kind == R.KIND_TEXT:
            texts[role.key] = [c.strip() for c in cells]
            continue
        if role.kind in (R.KIND_TIME, R.KIND_DATE, R.KIND_CLOCK, R.KIND_DAY):
            if role.kind == R.KIND_TIME:
                values, bad, used, ambiguous = parse_timestamps(cells, entry.get("format", "auto"), offset, missing)
                time_format = used
            elif role.kind == R.KIND_DATE:
                values, bad, used, ambiguous = parse_dates(cells, entry.get("format", "auto"), missing)
                time_format = f"{used} + time of day"
            elif role.kind == R.KIND_CLOCK:
                values, bad = parse_clock(cells, missing)
                ambiguous = False
            else:
                values, bad = parse_numbers(cells, comma, missing)
                ambiguous = False
                bad = bad | (np.isfinite(values) & ((values < 1) | (values != np.floor(values))))
            if ambiguous:
                warnings.append(f"{header}: day and month order could be either way; read as "
                                f"{_order_text(time_format)}. Check the time range.")
            reject(bad, lambda i, h=header, c=cells, lb=label: f"{h}: '{c[i].strip()}' is not a {lb.lower()}")
            time_parts[role.key] = values
            continue
        if role.kind == R.KIND_COORD:
            values, bad = parse_coordinates(cells, comma, missing, 90.0 if role.key == R.LAT else 180.0)
            reject(bad, lambda i, h=header, c=cells, lb=label: f"{h}: '{c[i].strip()}' is not a {lb.lower()}")
            key, unit = role.key, "deg"
        else:
            values, bad = parse_numbers(cells, comma, missing)
            reject(bad, lambda i, h=header, c=cells: f"{h}: '{c[i].strip()}' is not a number")
            unit_key = entry.get("unit") or role.default_unit
            unit = role.unit(unit_key) if role.units else None
            if unit is not None and unit.factor != 1.0:
                values = values * unit.factor
            if role.sign:
                values = values * R.sign_factor(role.key, entry.get("sign", ""))
            key = P.entry_channel_key(entry)
            unit = role.units[0].label if role.units else (entry.get("unit") or "")
            if role.key in (R.EASTING, R.NORTHING):
                unit = "m"
        columns[key] = values
        channels.append(Channel(key, P.entry_label(entry), role.key, unit, header))

    if time_parts:
        if R.TIME in time_parts:
            stamp = time_parts[R.TIME]
        else:
            if R.DATE in time_parts:
                base = time_parts[R.DATE]
            else:
                start = datetime.strptime(prof["start_date"].strip(), "%Y-%m-%d")
                base = _naive_seconds(start) + (time_parts[R.DAY] - 1) * 86400.0
            stamp = base + time_parts[R.CLOCK] - offset
        missing_time = ~np.isfinite(stamp)
        reject(missing_time, lambda i: "no time")
        columns[R.TIME] = stamp
        channels.insert(0, Channel(R.TIME, "Time", R.TIME, "UTC", _time_headers(entries)))
    reject(~np.isfinite(columns[R.KP]), lambda i: "no KP")

    known = {e["header"] for e in prof["columns"]}
    if prof.get("keep_extra", True):
        used_keys = set(columns)
        for header in headers:
            if header in known:
                continue
            cells = cells_of(index[header])
            values, bad = parse_numbers(cells, comma, missing)
            filled = int(np.isfinite(values).sum())
            if filled == 0 or filled < _EXTRA_NUMERIC_SHARE * (filled + int(bad.sum())):
                continue
            values[bad] = np.nan
            key = R.channel_key(R.EXTRA, header)
            while key in used_keys:
                key += "_"
            used_keys.add(key)
            unit = raw.units[index[header]].strip() if index[header] < len(raw.units) else ""
            columns[key] = values
            channels.append(Channel(key, header, R.EXTRA, unit, header, auto=True))

    keep = np.array([r is None for r in reasons], dtype=bool)
    rejected = [(raw.lines[i], reasons[i], raw.row_text(i)) for i in range(n) if reasons[i] is not None]
    for key in list(columns):
        columns[key] = np.ascontiguousarray(columns[key][keep], dtype=np.float64)
    for key in list(texts):
        texts[key] = [t for t, k in zip(texts[key], keep) if k]
    line = np.asarray(raw.lines, dtype=np.int64)[keep]
    parsed = ParsedData(columns, texts, channels, line, rejected, warnings, n, time_format)
    parsed.warnings.extend(_checks(parsed))
    return parsed


def _order_text(fmt: str) -> str:
    if "%d" in fmt and "%m" in fmt and fmt.index("%d") < fmt.index("%m"):
        return "day first"
    return "month first"


def _time_headers(entries: Sequence[Dict]) -> str:
    return " + ".join(e["header"] for e in entries
                      if R.role(e["role"]).kind in (R.KIND_TIME, R.KIND_DATE, R.KIND_CLOCK, R.KIND_DAY))


def _checks(parsed: ParsedData) -> List[str]:
    """Plain warnings about values that look wrong (nothing is changed)."""
    out: List[str] = []
    if not parsed.n:
        out.append("No rows could be read.")
        return out
    kp = parsed.columns[R.KP]
    kp_unit = next((c for c in parsed.channels if c.key == R.KP), None)
    finite_kp = kp[np.isfinite(kp)]
    if finite_kp.size and kp_unit is not None:
        if np.nanmax(np.abs(finite_kp)) > 5000:
            out.append("KP values exceed 5,000 km: are they in metres?")
    dol = parsed.columns.get(R.DOL)
    if dol is not None and np.isfinite(dol).any():
        negative = int((dol < 0).sum())
        if negative > 0.5 * np.isfinite(dol).sum():
            out.append(f"{negative:,} DOL values are negative: is the sign positive up?")
        if np.nanmax(dol) > 20:
            out.append("DOL values exceed 20 m: are they in cm or mm?")
    stamp = parsed.columns.get(R.TIME)
    if stamp is not None and stamp.size > 1:
        steps = np.diff(stamp)
        back = int((steps < 0).sum())
        if back:
            out.append(f"Time goes backwards {back:,} times (rows are kept in file order).")
        repeats = int((steps == 0).sum())
        if repeats:
            out.append(f"{repeats:,} rows repeat the previous timestamp.")
    return out


def read(content: bytes, name: str, profile: Dict, cancelled: Optional[Callable[[], bool]] = None
         ) -> Tuple[RawTable, ParsedData]:
    """Read and parse a whole delivery with ``profile``."""
    raw = read_table(content, name, P.normalise(profile)["table"])
    return raw, parse(raw, profile, cancelled)
