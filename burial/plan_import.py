"""Import an existing burial plan from a KP-range table (no Qt).

Contractor and client plans arrive as spreadsheets in every layout: one row
per KP range with a column saying what happens there ("Plough", "Skip",
"PLB", "Surface lay", "Y"/"N"...), sometimes a tool column, sometimes only
the burial ranges. This module turns such a table into plan events:

1. ``read_grid`` — raw rows from CSV/TSV/TXT or XLSX (no header guessing).
2. ``guess_header_row`` / ``guess_roles`` — suggest the header row and which
   column holds start KP, end KP, action, tool and notes.
3. ``distinct_values`` + ``guess_action`` / ``guess_tool`` — suggest how each
   value in the action/tool columns maps to Bury / Skip / Ignore and to a
   registered burial tool.
4. ``build_plan`` — validated, scope-clipped, merged burial ranges, the
   START/END events that encode them in the plan's travel direction, and
   per-range tool/notes to stamp on the derived sections.

Everything is plain data so the wizard, tests and future importers share it.
"""
from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

from . import schema

ROLE_IGNORE = ""
ROLE_START = "start_kp"
ROLE_END = "end_kp"
ROLE_ACTION = "action"
ROLE_TOOL = "tool"
ROLE_NOTES = "notes"
ROLES = [ROLE_IGNORE, ROLE_START, ROLE_END, ROLE_ACTION, ROLE_TOOL, ROLE_NOTES]
ROLE_LABELS = {
    ROLE_IGNORE: "(ignore)",
    ROLE_START: "Start KP",
    ROLE_END: "End KP",
    ROLE_ACTION: "Action (bury / skip)",
    ROLE_TOOL: "Tool",
    ROLE_NOTES: "Notes",
}

ACTION_BURY = "bury"
ACTION_SKIP = "skip"
ACTION_IGNORE = "ignore"
ACTIONS = [ACTION_BURY, ACTION_SKIP, ACTION_IGNORE]
ACTION_LABELS = {ACTION_BURY: "Bury / tool section", ACTION_SKIP: "Skip (no burial)",
                 ACTION_IGNORE: "Ignore row"}

MODE_REPLACE = "replace"
MODE_MERGE = "merge"

# Travel-direction KP tolerance shared with the event validator.
_KP_TOL = 5e-7
# Header-row search depth and the preview row cap for very large files.
_HEADER_SCAN_ROWS = 15
MAX_ROWS = 20000

_ROLE_SYNONYMS: Dict[str, Tuple[str, ...]] = {
    ROLE_START: ("start kp", "kp start", "from kp", "kp from", "start", "from",
                 "kp1", "begin kp", "start chainage", "from chainage", "kp_start",
                 "start_kp", "kp in"),
    ROLE_END: ("end kp", "kp end", "to kp", "kp to", "end", "to", "kp2",
               "end chainage", "to chainage", "kp_end", "end_kp", "kp out",
               "finish kp", "stop kp"),
    ROLE_ACTION: ("action", "activity", "operation", "burial", "bury", "burial method",
                  "method", "plan", "burial plan", "protection",
                  "installation", "lay method", "treatment", "section type"),
    ROLE_TOOL: ("tool", "burial tool", "equipment", "asset", "spread", "vessel tool",
                "machine"),
    ROLE_NOTES: ("notes", "note", "remarks", "remark", "comment", "comments",
                 "description", "reason"),
}

# Value keywords (casefolded, word-ish match). Skip wins over bury so
# "no burial" / "not buried" / "skip - plough recovered" read as skips.
_SKIP_WORDS = ("skip", "no burial", "not buried", "unburied", "surface lay",
               "surface laid", "no bury", "none", "exclusion", "excluded",
               "transit", "recover", "lift", "n/a", "na", "no",
               "n", "false", "0", "-", "x")
_BURY_WORDS = ("bury", "burial", "buried", "plough", "plow", "pldn", "trench",
               "trencher", "jet", "jetting", "rov", "plb", "post lay", "post-lay",
               "mfe", "mass flow", "cutter", "pre-cut", "precut", "yes", "y",
               "true", "1", "simultaneous", "slb", "inspection", "inspect")

# Tool keyword → registered tool_type (method) guess.
_TOOL_TYPE_WORDS = (
    (("plough", "plow", "pldn", "plup"), schema.METHOD_PLOUGH),
    (("mfe", "mass flow"), schema.METHOD_MFE),
    (("trench", "jet", "rov", "plb", "post lay", "post-lay", "cutter"),
     schema.METHOD_TRENCHER),
    (("inspection", "inspect", "survey"), schema.METHOD_INSPECTION),
)


# ---------------------------------------------------------------- reading
def read_grid(path: str, sheet: Optional[str] = None,
              max_rows: int = MAX_ROWS) -> Tuple[List[List[str]], List[str]]:
    """``(rows, sheet_names)``: every non-empty row as stripped strings.

    Rows are padded to a common width. No header detection happens here —
    the wizard lets the user pick (or confirm) the header row.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        rows, names = _read_xlsx(path, sheet, max_rows)
    else:
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            rows, names = parse_delimited(handle.read(), max_rows), []
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows], names


def parse_delimited(text: str, max_rows: int = MAX_ROWS) -> List[List[str]]:
    lines = [ln for ln in text.splitlines()
             if ln.strip() and not ln.lstrip().startswith("#")]
    if not lines:
        return []
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:30]), delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO("\n".join(lines)), dialect):
        cells = [cell.strip() for cell in row]
        if any(cells):
            rows.append(cells)
        if len(rows) >= max_rows:
            break
    return rows


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return str(value).strip()


def _read_xlsx(path: str, sheet: Optional[str], max_rows: int
               ) -> Tuple[List[List[str]], List[str]]:
    try:
        import openpyxl
    except Exception as exc:  # pragma: no cover — environment specific
        raise ValueError(
            "openpyxl is required to read Excel files but could not be "
            f"imported. Ensure the plugin's lib/ folder is present. ({exc})")
    if getattr(openpyxl, "LXML", False):
        raise ValueError(
            "Excel support was loaded with an unsafe native XML backend. "
            "Restart QGIS once to activate the safe workbook reader.")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(book.sheetnames)
        ws = book[sheet] if sheet and sheet in names else book[names[0]]
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            if any(cells):
                while cells and not cells[-1]:
                    cells.pop()
                rows.append(cells)
            if len(rows) >= max_rows:
                break
    finally:
        try:
            book.close()
        except Exception:
            pass
    return rows, names


# ---------------------------------------------------------------- guessing
def parse_kp(text, unit_factor: float = 1.0) -> Optional[float]:
    """KP in km from a cell: '12.345', '12,345' (decimal comma), 'KP 12.3',
    '12+345' (km+m chainage). ``unit_factor`` converts other units to km."""
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return None
    s = re.sub(r"(?i)^\s*r?kp\s*", "", s).replace(" ", "")
    match = re.fullmatch(r"(-?\d+)\+(\d+(?:\.\d+)?)", s)
    if match:
        return float(match.group(1)) + float(match.group(2)) / 1000.0
    if s.count(",") == 1 and "." not in s:
        s = s.replace(",", ".")
    s = re.sub(r"(?i)(km|m)$", "", s)
    try:
        return float(s) * unit_factor
    except ValueError:
        return None


def _norm(text: str) -> str:
    out = (text or "").strip().casefold().replace("_", " ")
    out = re.sub(r"[\(\[][^\)\]]*[\)\]]", " ", out)
    return " ".join(out.replace(":", " ").split())


def guess_header_row(rows: Sequence[Sequence[str]]) -> int:
    """Index of the header row: the last mostly-text row before KP data starts.

    Returns -1 when the first row already holds numeric KPs (no header).
    """
    def numeric_pair(row) -> bool:
        return sum(1 for c in row if parse_kp(c) is not None) >= 2

    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        if numeric_pair(row):
            # Walk back to the nearest non-numeric row with ≥2 texts.
            for j in range(i - 1, -1, -1):
                if sum(1 for c in rows[j] if c) >= 2 and not numeric_pair(rows[j]):
                    return j
            return -1
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        if sum(1 for c in row if c) >= 2:
            return i
    return -1


def guess_roles(headers: Sequence[str],
                data_rows: Sequence[Sequence[str]] = ()) -> List[str]:
    """One role per column from the header text, falling back to content
    (first two mostly-numeric columns → start/end KP)."""
    normed = [_norm(h) for h in headers]
    roles = [ROLE_IGNORE] * len(headers)
    for exact in (True, False):
        for role, synonyms in _ROLE_SYNONYMS.items():
            if role in roles:
                continue
            for syn in sorted(synonyms, key=len, reverse=True):
                hit = next((i for i, h in enumerate(normed)
                            if roles[i] == ROLE_IGNORE and h
                            and ((h == syn) if exact else (len(syn) >= 4 and syn in h))),
                           None)
                if hit is not None:
                    roles[hit] = role
                    break
    if ROLE_START not in roles or ROLE_END not in roles:
        sample = list(data_rows[:50])
        numeric = [i for i in range(len(headers)) if roles[i] == ROLE_IGNORE and sample
                   and sum(1 for r in sample if i < len(r) and parse_kp(r[i]) is not None)
                   >= 0.8 * len(sample)]
        for role in (ROLE_START, ROLE_END):
            if role not in roles and numeric:
                roles[numeric.pop(0)] = role
    return roles


def _has_word(text: str, word: str) -> bool:
    return re.search(r"(?<![a-z0-9])" + re.escape(word) + r"(?![a-z0-9])", text) is not None


def _norm_value(text: str) -> str:
    """Cell value for keyword matching (brackets kept: '(skip)' matters)."""
    return " ".join((text or "").strip().casefold().replace("_", " ").split())


def guess_action(value: str) -> str:
    text = _norm_value(value)
    if not text:
        return ACTION_IGNORE
    if any(_has_word(text, w) for w in _SKIP_WORDS):
        return ACTION_SKIP
    if any(_has_word(text, w) for w in _BURY_WORDS):
        return ACTION_BURY
    return ACTION_IGNORE


def guess_tool(value: str, tools: Sequence[Dict]) -> str:
    """Registered ``tool_id`` for a tool/action cell ("" = plan default)."""
    text = _norm_value(value)
    if not text:
        return ""
    for tool in tools or []:
        name = _norm_value(tool.get("name") or "")
        if name and (name == text or (len(name) >= 3 and name in text)
                     or (len(text) >= 3 and text in name)):
            return str(tool.get("tool_id") or "")
    for words, tool_type in _TOOL_TYPE_WORDS:
        if any(_has_word(text, w) for w in words):
            matches = [t for t in tools or []
                       if schema.normalise_method(t.get("tool_type") or "") == tool_type]
            if len(matches) == 1:
                return str(matches[0].get("tool_id") or "")
    return ""


def distinct_values(rows: Sequence[Sequence[str]], column: int) -> List[Tuple[str, int]]:
    """``[(value, count)]`` in first-seen order for one column."""
    counts: Dict[str, int] = {}
    for row in rows:
        value = row[column].strip() if 0 <= column < len(row) else ""
        counts[value] = counts.get(value, 0) + 1
    return list(counts.items())


# ---------------------------------------------------------------- building
@dataclass
class ImportRange:
    row: int                 # 1-based row number in the source table
    start_kp: float
    end_kp: float
    action: str
    tool_id: str = ""
    notes: str = ""


@dataclass
class ImportResult:
    ranges: List[ImportRange] = field(default_factory=list)   # usable rows, KP order
    burial: List[ImportRange] = field(default_factory=list)   # merged burial ranges
    events: List[Dict] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    skipped_rows: int = 0
    transitions: int = 0      # continuous burial handing over to another tool

    @property
    def burial_km(self) -> float:
        return sum(r.end_kp - r.start_kp for r in self.burial)


@dataclass
class ImportSpec:
    roles: List[str]
    action_map: Dict[str, str] = field(default_factory=dict)   # value → ACTION_*
    tool_map: Dict[str, str] = field(default_factory=dict)     # value → tool_id
    kp_factor: float = 1.0          # multiply cells to get km (0.001 for metres)
    default_action: str = ACTION_BURY  # rows with no action column

    def column(self, role: str) -> int:
        return self.roles.index(role) if role in self.roles else -1


def _new_event(kp: float, event_type: str, note: str = "") -> Dict:
    return {
        "event_id": schema.new_id(), "plan_id": "", "generation_id": "",
        "seq": 0, "event_type": event_type, "kp": round(kp, 6), "end_kp": None,
        "lat": None, "lon": None, "depth_m": None,
        "source": schema.EVENT_SOURCE_IMPORT,
        "status": schema.EVENT_STATUS_CANDIDATE, "locked": 0, "notes": note or "",
    }


def ranges_to_events(burial: Sequence[ImportRange], direction: int) -> List[Dict]:
    """START/END pairs in the plan's travel direction (−1: START at the high KP)."""
    events = []
    for rng in burial:
        first, last = (rng.start_kp, rng.end_kp) if direction >= 0 else (rng.end_kp, rng.start_kp)
        events.append(_new_event(first, schema.EVENT_BURIAL_START))
        events.append(_new_event(last, schema.EVENT_BURIAL_END))
    return events


def _fmt(kp: float) -> str:
    return schema.format_kp(kp)


def build_plan(data_rows: Sequence[Sequence[str]], spec: ImportSpec,
               scope: Tuple[float, float], direction: int = 1,
               first_row_number: int = 2, tool_names: Optional[Dict[str, str]] = None,
               ) -> ImportResult:
    """Validate mapped rows and turn the burial ranges into plan events.

    * Rows missing a KP, or mapped to *Ignore*, are skipped (counted).
    * Ranges are clipped to the plan scope; ranges wholly outside are dropped.
    * Overlapping rows are errors (the plan cannot say both bury and skip).
    * Touching burial ranges with the same tool merge into one section;
      with different tools they meet at a tool transition (END and START
      at the same KP, e.g. PLUP / Start PLB).
    * Gaps between burial ranges become skip sections when events are
      derived, so explicit skip rows only document intent.
    """
    result = ImportResult()
    tool_names = tool_names or {}
    c_start, c_end = spec.column(ROLE_START), spec.column(ROLE_END)
    c_action, c_tool, c_notes = (spec.column(ROLE_ACTION), spec.column(ROLE_TOOL),
                                 spec.column(ROLE_NOTES))
    if c_start < 0 or c_end < 0:
        result.errors.append("Map a Start KP and an End KP column.")
        return result
    scope_lo, scope_hi = sorted(float(v) for v in scope)

    def cell(row, col):
        return row[col].strip() if 0 <= col < len(row) else ""

    clipped = 0
    outside = 0
    for offset, row in enumerate(data_rows):
        number = first_row_number + offset
        if not any((c or "").strip() for c in row):
            continue
        action = spec.action_map.get(cell(row, c_action), ACTION_IGNORE) \
            if c_action >= 0 else spec.default_action
        if action == ACTION_IGNORE:
            result.skipped_rows += 1
            continue
        a, b = parse_kp(cell(row, c_start), spec.kp_factor), parse_kp(cell(row, c_end), spec.kp_factor)
        if a is None or b is None:
            result.skipped_rows += 1
            result.warnings.append(f"Row {number}: no numeric start/end KP — skipped.")
            continue
        lo, hi = min(a, b), max(a, b)
        if hi - lo <= _KP_TOL:
            result.skipped_rows += 1
            result.warnings.append(f"Row {number}: zero-length range at KP {_fmt(lo)} — skipped.")
            continue
        if hi <= scope_lo + _KP_TOL or lo >= scope_hi - _KP_TOL:
            outside += 1
            continue
        if lo < scope_lo - _KP_TOL or hi > scope_hi + _KP_TOL:
            clipped += 1
            lo, hi = max(lo, scope_lo), min(hi, scope_hi)
        tool_value = cell(row, c_tool) if c_tool >= 0 else ""
        tool_id = spec.tool_map.get(tool_value, "") if tool_value else ""
        if not tool_id and c_tool < 0 and c_action >= 0:
            # No tool column: an action value such as "Plough" may name one.
            tool_id = spec.tool_map.get(cell(row, c_action), "")
        result.ranges.append(ImportRange(number, lo, hi, action, tool_id, cell(row, c_notes)))

    if outside:
        result.warnings.append(
            f"{outside} range(s) lie wholly outside the plan scope "
            f"KP {_fmt(scope_lo)}–{_fmt(scope_hi)} and were left out.")
    if clipped:
        result.warnings.append(f"{clipped} range(s) were clipped to the plan scope.")

    result.ranges.sort(key=lambda r: (r.start_kp, r.end_kp))
    for prev, cur in zip(result.ranges, result.ranges[1:]):
        if cur.start_kp < prev.end_kp - _KP_TOL:
            result.errors.append(
                f"Rows {prev.row} and {cur.row} overlap (KP {_fmt(cur.start_kp)}–"
                f"{_fmt(min(prev.end_kp, cur.end_kp))}). Fix the table or ignore one of them.")
    if result.errors:
        return result

    merged: List[ImportRange] = []
    for rng in (r for r in result.ranges if r.action == ACTION_BURY):
        last = merged[-1] if merged else None
        touching = last is not None and rng.start_kp <= last.end_kp + _KP_TOL
        if touching and (rng.tool_id or "") != (last.tool_id or ""):
            # Continuous burial with another tool: a tool transition (e.g.
            # PLUP and Start PLB at one KP), not a merge.
            result.transitions += 1
            merged.append(ImportRange(rng.row, rng.start_kp, rng.end_kp, ACTION_BURY,
                                      rng.tool_id, rng.notes))
        elif touching:
            notes = [n for n in (last.notes, rng.notes) if n]
            merged[-1] = ImportRange(last.row, last.start_kp, max(last.end_kp, rng.end_kp),
                                     ACTION_BURY, last.tool_id or rng.tool_id,
                                     "; ".join(dict.fromkeys(notes)))
        else:
            merged.append(ImportRange(rng.row, rng.start_kp, rng.end_kp, ACTION_BURY,
                                      rng.tool_id, rng.notes))
    result.burial = merged
    if not merged:
        result.warnings.append("No burial ranges: the imported plan is all skip.")
    result.events = ranges_to_events(merged, direction)
    return result


def _existing_burial(events: Sequence[Dict], direction: int,
                     scope: Tuple[float, float]) -> List[Tuple[float, float]]:
    from .events import burial_pairs
    lo_scope, hi_scope = sorted(float(v) for v in scope)
    out = []
    for start, end in burial_pairs(list(events), direction):
        a = float(start.get("kp"))
        b = float(end.get("kp")) if end is not None else (hi_scope if direction >= 0 else lo_scope)
        lo, hi = min(a, b), max(a, b)
        if hi - lo > _KP_TOL:
            out.append((lo, hi))
    return out


def _subtract(intervals, cuts):
    out = []
    for lo, hi in intervals:
        pieces = [(lo, hi)]
        for c_lo, c_hi in cuts:
            nxt = []
            for p_lo, p_hi in pieces:
                if c_hi <= p_lo + _KP_TOL or c_lo >= p_hi - _KP_TOL:
                    nxt.append((p_lo, p_hi))
                    continue
                if c_lo > p_lo + _KP_TOL:
                    nxt.append((p_lo, c_lo))
                if c_hi < p_hi - _KP_TOL:
                    nxt.append((c_hi, p_hi))
            pieces = nxt
        out.extend(pieces)
    return out


def plan_events(existing: Sequence[Dict], result: ImportResult, mode: str,
                direction: int, scope: Tuple[float, float],
                sections: Sequence[Dict] = ()) -> Tuple[List[Dict], List[Dict]]:
    """``(events, dropped)`` for committing ``result`` in ``mode``.

    * ``MODE_REPLACE``: the imported burial ranges become the whole plan.
    * ``MODE_MERGE`` (overlay): inside the KP ranges the file covers (bury
      *and* skip rows) the file wins; elsewhere the existing plan stays.

    Existing START/END events are reused where a final boundary lands on
    one, keeping their IDs, lock/confirm state and notes. ``dropped`` lists
    existing events that no longer exist (so the UI can warn about locked
    or confirmed ones).
    """
    existing = [dict(e) for e in existing]
    boundary = [e for e in existing
                if e.get("event_type") in (schema.EVENT_BURIAL_START, schema.EVENT_BURIAL_END)]
    others = [e for e in existing if e not in boundary]
    if mode == MODE_MERGE:
        covered = [(r.start_kp, r.end_kp) for r in result.ranges]
        kept = _subtract(_existing_burial(boundary, direction, scope), covered)
        # (lo, hi, source, tool): existing pieces never merge with each other
        # (touching ones are existing tool transitions); an imported range
        # merges into a touching existing piece when it names no tool or the
        # same tool as the existing section there (``sections``).
        def existing_tool(lo, hi):
            mid = (lo + hi) / 2.0
            hit = next((s for s in sections if s.get("kind") == schema.SECTION_BURIAL
                        and float(s.get("start_kp") or 0.0) <= mid
                        <= float(s.get("end_kp") or 0.0)), None)
            return (hit or {}).get("tool_id") or ""

        final = sorted([(lo, hi, "kept", existing_tool(lo, hi)) for lo, hi in kept]
                       + [(r.start_kp, r.end_kp, "file", r.tool_id or "") for r in result.burial])
        merged: List[List] = []
        for lo, hi, source, tool in final:
            last = merged[-1] if merged else None
            joinable = (last is not None and lo <= last[1] + _KP_TOL
                        and source != last[2]
                        and (tool == last[3] or not (tool if source == "file" else last[3])))
            if joinable:
                last[1] = max(last[1], hi)
            else:
                merged.append([lo, hi, source, tool])
        intervals = [(m[0], m[1]) for m in merged]
    else:
        intervals = [(r.start_kp, r.end_kp) for r in result.burial]
    wanted = ranges_to_events([ImportRange(0, lo, hi, ACTION_BURY) for lo, hi in intervals],
                              direction)
    pool = list(boundary)
    events = []
    for event in wanted:
        match = next((e for e in pool if e.get("event_type") == event["event_type"]
                      and abs(float(e.get("kp")) - event["kp"]) <= 1e-6), None)
        if match is not None:
            pool.remove(match)
            events.append(match)
        else:
            events.append(event)
    return events + others, pool


def section_updates(sections: Sequence[Dict], burial: Sequence[ImportRange],
                    skip_notes: Sequence[ImportRange] = ()) -> Dict[str, Dict]:
    """``{section_id: updates}`` giving derived sections the imported tool/notes.

    Burial sections match imported burial ranges by KP; skip sections take
    the notes of skip rows lying inside them.
    """
    out: Dict[str, Dict] = {}
    tol = 1e-5
    for section in sections:
        try:
            lo, hi = sorted((float(section.get("start_kp")), float(section.get("end_kp"))))
        except (TypeError, ValueError):
            continue
        sid = str(section.get("section_id") or "")
        if section.get("kind") == schema.SECTION_BURIAL:
            rng = next((r for r in burial if abs(r.start_kp - lo) <= tol
                        and abs(r.end_kp - hi) <= tol), None)
            if rng is None:
                continue
            updates = {}
            if rng.tool_id:
                updates["tool_id"] = rng.tool_id
            if rng.notes:
                updates["notes"] = rng.notes
            if updates:
                out[sid] = updates
        elif section.get("kind") == schema.SECTION_SKIP:
            notes = [r.notes for r in skip_notes
                     if r.notes and r.start_kp >= lo - tol and r.end_kp <= hi + tol]
            if notes:
                out[sid] = {"notes": "; ".join(dict.fromkeys(notes))}
    return out
