# -*- coding: utf-8 -*-
"""Spreadsheet files for Route Comments (no QGIS): read CSV/TSV/XLSX, write CSV/XLSX.

Reading reuses the Burial Planner's table reader (any delimiter, any
sheet, rows as text). XLSX uses the bundled openpyxl, imported only when
an Excel file is read or written.
"""

from __future__ import annotations

import csv
import os
from typing import List, Optional, Sequence, Tuple

from ..burial.plan_import import read_grid

MAX_ROWS = 50000


def read_table(path: str, sheet: Optional[str] = None) -> Tuple[List[List[str]], List[str]]:
    """``(rows, sheet names)`` of a CSV/TSV/TXT or XLSX file; rows as text."""
    return read_grid(path, sheet, MAX_ROWS)


def _cell(value):
    return "" if value is None else value


def write_csv(path: str, headers: Sequence[str], rows: Sequence[Sequence]) -> None:
    """UTF-8 with a byte-order mark so spreadsheet programs read accents and dashes."""
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(list(headers))
        for row in rows:
            writer.writerow([_cell(v) for v in row])


def write_xlsx(path: str, headers: Sequence[str], rows: Sequence[Sequence], about: Sequence[Tuple[str, str]] = ()
               ) -> None:
    """A *Comments* sheet (frozen header, filters) and an *About* sheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except Exception as exc:  # pragma: no cover - environment specific
        raise ValueError("Excel export needs the bundled openpyxl library, which could not be loaded "
                         f"({exc}). Export to CSV instead.")
    book = Workbook()
    sheet = book.active
    sheet.title = "Comments"
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_cell(v) for v in row])
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    widths = {"Comment": 60, "Reply text (export only)": 50, "Geometry (WKT, EPSG:4326)": 40}
    for index, header in enumerate(headers, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = widths.get(header, max(10, min(28, len(str(header)) + 4)))
        if header in widths:
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    info = book.create_sheet("About")
    for key, value in about:
        info.append([key, value])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 110
    for cell in info["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(path)


def write_table(path: str, headers: Sequence[str], rows: Sequence[Sequence], about: Sequence[Tuple[str, str]] = ()
                ) -> None:
    """CSV or XLSX by the file extension."""
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm"):
        write_xlsx(path, headers, rows, about)
    else:
        write_csv(path, headers, rows)
